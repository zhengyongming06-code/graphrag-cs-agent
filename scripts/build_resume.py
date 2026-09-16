# -*- coding: utf-8 -*-
"""Clean one-page resume: keep core amplify + metrics, reduce visual noise."""
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

OUT_DOCX = Path(r"d:\Desktop\郑咏明-AI-Agent实习简历.docx")


def set_run_font(run, name="微软雅黑", size=10.5, bold=False, color=None):
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    if color is not None:
        run.font.color.rgb = color


def add_center(doc, text, size=11, bold=False, after=2, color=None):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.05
    set_run_font(p.add_run(text), size=size, bold=bold, color=color)


def add_section(doc, title):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.05
    set_run_font(p.add_run(title), size=12, bold=True, color=RGBColor(0x1F, 0x3A, 0x5F))
    pPr = p._p.get_or_add_pPr()
    pBdr = pPr.makeelement(qn("w:pBdr"), {})
    bottom = pBdr.makeelement(
        qn("w:bottom"),
        {qn("w:val"): "single", qn("w:sz"): "10", qn("w:space"): "1", qn("w:color"): "94A3B8"},
    )
    pBdr.append(bottom)
    pPr.append(pBdr)


def add_line(doc, text, bold=False, size=10.5, before=1, after=2, color=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.12
    set_run_font(p.add_run(text), size=size, bold=bold, color=color)


def add_bullet(doc, text, size=10.5):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.left_indent = Cm(0.4)
    p.clear()
    set_run_font(p.add_run(text), size=size)


def main():
    doc = Document()
    for section in doc.sections:
        section.top_margin = Cm(1.3)
        section.bottom_margin = Cm(1.3)
        section.left_margin = Cm(1.6)
        section.right_margin = Cm(1.6)

    # Header — simple
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(2)
    set_run_font(p.add_run("郑咏明"), size=20, bold=True)

    add_center(doc, "求职意向：AI Agent / 大模型应用开发实习生", size=11, bold=True, after=2)
    add_center(
        doc,
        "18096270420　|　Jas-4ever@outlook.com　|　github.com/zhengyongming06-code",
        size=9.5,
        after=2,
        color=RGBColor(0x47, 0x55, 0x69),
    )

    # Education
    add_section(doc, "教育背景")
    add_line(doc, "天津城建大学　|　计算机科学与技术　|　本科在读　|　2025.09 — 2029.06", bold=True, size=10.5)

    # Skills — short lines, no cram
    add_section(doc, "专业技能")
    add_bullet(
        doc,
        "Agent：LangGraph 状态机、ReAct、Tool Calling、任务拆解；证据不足时转人工工单（Human-in-the-loop）。",
    )
    add_bullet(
        doc,
        "RAG / GraphRAG：Neo4j；Vector + Lexical + 1-hop Graph 混合召回；RRF（k=60）融合；切分与 Hit@3 评测。",
    )
    add_bullet(
        doc,
        "工程：Python、FastAPI、Docker Compose、DeepSeek / OpenAI 兼容 API、Git；Citation 与 Tool Trace。",
    )

    # Projects
    add_section(doc, "项目经历")

    # ---- Project 1 ----
    add_line(doc, "知识库 GraphRAG 客服　　2026.09", bold=True, size=11, before=2, after=1)
    add_line(
        doc,
        "Python / FastAPI / LangGraph / Neo4j / DeepSeek / Docker Compose",
        size=9.5,
        before=0,
        after=1,
        color=RGBColor(0x47, 0x55, 0x69),
    )
    add_line(
        doc,
        "仓库：https://github.com/zhengyongming06-code/graphrag-cs-agent",
        size=9.5,
        before=0,
        after=2,
        color=RGBColor(0x47, 0x55, 0x69),
    )
    add_line(
        doc,
        "自建知识库客服：文档切分入库、混合检索、引用作答，证据不够则转人工。",
        size=10.5,
        before=0,
        after=3,
    )
    add_bullet(
        doc,
        "图谱：Document–Chunk–Entity；共现实体 1-hop 扩展。",
    )
    add_bullet(
        doc,
        "检索：向量 + 关键词 + 图谱，RRF 融合。自建问答 Hit@3 相对纯向量有提升。",
    )
    add_bullet(
        doc,
        "对话：LangGraph 走记忆、检索、门控、作答或建工单。",
    )
    add_bullet(
        doc,
        "Docker 拉起 Neo4j；页面可对话、补文档、对比两种检索。",
    )

    # ---- Project 2 ----
    add_line(doc, "NovaPulse DeepResearch 多智能体调研中台（独立研发）　　2026.09", bold=True, size=11, before=8, after=1)
    add_line(
        doc,
        "Python / FastAPI / LangGraph / DeepSeek / Trafilatura / Map-Reduce / Docker",
        size=9.5,
        before=0,
        after=2,
        color=RGBColor(0x47, 0x55, 0x69),
    )
    add_line(
        doc,
        "仓库：https://github.com/zhengyongming06-code/novapulse-deepresearch",
        size=9.5,
        before=0,
        after=2,
        color=RGBColor(0x47, 0x55, 0x69),
    )
    add_line(
        doc,
        "面向竞品/技术调研，做成可追踪的多智能体流水线："
        "计划 → 检索 → 正文抽取 → Map 摘要 → Reduce 成稿 → 红队审稿，区别于低代码拖拽编排。",
        size=10.5,
        before=0,
        after=3,
    )
    add_bullet(
        doc,
        "编排：LangGraph 命名节点状态机（plan/scout/harvest/distill/compose/redteam），"
        "SSE 回放每步，单步失败可降级到本地来源。",
    )
    add_bullet(
        doc,
        "取证：Bing/DuckDuckGo（Tavily 可选）；Trafilatura 去导航后再进模型；"
        "引用账本把结论编号绑定 URL。",
    )
    add_bullet(
        doc,
        "成稿：多源 Map-Reduce 控制上下文；红队节点检查无引用断言并打分；报告 md/json 落盘。",
    )

    doc.save(OUT_DOCX)
    print(f"DOCX={OUT_DOCX}")


if __name__ == "__main__":
    main()
